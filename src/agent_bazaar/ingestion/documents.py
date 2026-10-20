"""Deterministic PDF, DOCX, and plain-text extraction with chunking."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Literal

from docx import Document as DocxDocument
from pypdf import PdfReader

from agent_bazaar.config import IngestionSettings
from agent_bazaar.ingestion.models import ChunkInfo, NormalizedSpec
from agent_bazaar.protocols import IngestionError


def _tool_id(source: Path) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", source.stem.lower())
    slug = re.sub(r"-+", "-", slug).strip("-")
    return slug or "tool"


def _first_non_empty_line(text: str) -> str:
    for line in text.splitlines():
        stripped = line.strip()
        if stripped:
            return stripped
    return ""


def chunk_text(text: str, size: int, overlap: int, source: str = "") -> list[ChunkInfo]:
    if not text:
        return []
    if size <= 0:
        raise ValueError("chunk size must be positive")
    effective_overlap = max(0, min(overlap, size - 1))
    step = size - effective_overlap
    chunks: list[ChunkInfo] = []
    start = 0
    index = 1
    while start < len(text):
        end = min(start + size, len(text))
        chunks.append(
            ChunkInfo(
                chunk_id=f"chunk-{index:04d}",
                text=text[start:end],
                source=source,
                start=start,
                end=end,
            )
        )
        if end >= len(text):
            break
        start += step
        index += 1
    return chunks


def _pdf_title(reader: PdfReader) -> str:
    try:
        metadata = reader.metadata
    except Exception:
        return ""
    if not metadata:
        return ""
    title: Any = metadata.get("/Title")
    if not title:
        title = getattr(metadata, "title", None)
    return str(title).strip() if title else ""


def _extract_pdf(source: Path) -> tuple[str, str, int]:
    try:
        reader = PdfReader(str(source))
        pages = list(reader.pages)
        texts: list[str] = []
        for page in pages:
            extracted = page.extract_text()
            if extracted is not None:
                texts.append(extracted)
        title = _pdf_title(reader)
        page_count = len(pages)
    except Exception as exc:
        raise IngestionError(
            "parse_error",
            f"Failed to parse PDF: {exc}",
            source=str(source),
            details={"exception": type(exc).__name__},
        ) from exc
    return "\n".join(texts), title, page_count


def _extract_docx(source: Path) -> tuple[str, str]:
    try:
        document = DocxDocument(str(source))
        parts: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text
            if text.strip():
                parts.append(text)
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                line = " | ".join(cells).strip()
                if line:
                    parts.append(line)
        title = str(document.core_properties.title or "").strip()
    except Exception as exc:
        raise IngestionError(
            "parse_error",
            f"Failed to parse DOCX: {exc}",
            source=str(source),
            details={"exception": type(exc).__name__},
        ) from exc
    return "\n".join(parts), title


def _extract_text(source: Path) -> str:
    try:
        data = source.read_bytes()
    except OSError as exc:
        raise IngestionError(
            "read_error",
            f"Could not read text file: {exc}",
            source=str(source),
            details={"exception": type(exc).__name__},
        ) from exc
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def parse_document(
    source: Path,
    kind: Literal["pdf", "docx", "text"],
    settings: IngestionSettings,
) -> NormalizedSpec:
    if not source.is_file():
        raise IngestionError(
            "file_not_found",
            f"Source file not found: {source}",
            source=str(source),
        )
    metadata: dict[str, Any] = {}
    if kind == "pdf":
        text, core_title, page_count = _extract_pdf(source)
        metadata["page_count"] = page_count
    elif kind == "docx":
        text, core_title = _extract_docx(source)
    elif kind == "text":
        text = _extract_text(source)
        core_title = ""
    else:
        raise IngestionError(
            "unsupported_format",
            f"Unsupported document kind: {kind}",
            source=str(source),
        )
    title = core_title or _first_non_empty_line(text)
    metadata["title"] = title
    metadata["source"] = str(source)
    metadata["spec_format"] = kind
    chunks = chunk_text(
        text,
        settings.chunk_size_chars,
        settings.chunk_overlap_chars,
        source=str(source),
    )
    return NormalizedSpec(
        tool_id=_tool_id(source),
        source=str(source),
        spec_format=kind,
        title=title,
        chunks=chunks,
        metadata=metadata,
    )
