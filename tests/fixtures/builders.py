"""Fixture builders for DOCX and PDF ingestion tests."""

from __future__ import annotations

from pathlib import Path

from docx import Document


def write_sample_docx(path: Path) -> Path:
    document = Document()
    document.add_paragraph("Agent Bazaar sample document.")
    document.add_paragraph("This document is used to verify DOCX extraction.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "alpha"
    table.cell(0, 1).text = "beta"
    table.cell(1, 0).text = "gamma"
    table.cell(1, 1).text = "delta"
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


def write_sample_pdf(path: Path, text: str = "hello ingestion world") -> Path:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT\n/F1 18 Tf\n72 720 Td\n({escaped}) Tj\nET\n".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"endstream",
    ]
    payload = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for index, body in enumerate(objects, start=1):
        offsets.append(len(payload))
        payload.extend(f"{index} 0 obj\n".encode("ascii"))
        payload.extend(body)
        payload.extend(b"\nendobj\n")
    xref_offset = len(payload)
    payload.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    payload.extend(b"0000000000 65535 f \n")
    for offset in offsets:
        payload.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    payload.extend(b"trailer\n")
    payload.extend(f"<< /Size {len(objects) + 1} /Root 1 0 R >>\n".encode("ascii"))
    payload.extend(b"startxref\n")
    payload.extend(f"{xref_offset}\n".encode("ascii"))
    payload.extend(b"%%EOF\n")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(bytes(payload))
    return path
